#!/usr/bin/env python
import os
#import pycel
import openpyxl
from openpyxl import load_workbook
from pycel import ExcelCompiler
import traceback
import xlrd
import pandas as pd
file1="../512065202_p163_p419.xlsx"
file2="../512065202_p12157_p120.xlsx"

#try:
#    excel = ExcelCompiler(filename=file)
#except Exception as e:
#    traceback.print_exc()
#try:
#    excel=xlrd.open_workbook(file)
#    print("managed to open with exlrd")
#except Exception as e:
#    traceback.print_exc()   
try:
    wb = load_workbook(file1)
    #df = pd.read_excel(file) 
    print("managed to open file1")
    wb = load_workbook(file2)
    print("managed to open file2")
except Exception as e:
    traceback.print_exc()   
        